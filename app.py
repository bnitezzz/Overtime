import io
import os
import re
import unicodedata
from datetime import date, datetime, time, timedelta
import uuid

import openpyxl
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Font, Alignment
import pandas as pd
import streamlit as st

st.set_page_config(page_title="Gestión horas extras", page_icon="⏱️", layout="wide")

FERIADOS_FIJOS = {(1, 1), (4, 19), (5, 1), (6, 24), (7, 5), (7, 24), (10, 12), (12, 24), (12, 25), (12, 31)}
NOMBRE_HOJA_PREFERIDA = "Horas Extras"
FILA_INICIO_DATOS = 13
FILA_FIN_DATOS = 27
MAX_JORNADAS_PLANTILLA = FILA_FIN_DATOS - FILA_INICIO_DATOS + 1
PLANTILLA_DEFAULT_PATH = "plantilla.xlsx"
COLOR_ROJO = "FFFF0000"


def normalizar_texto(valor) -> str:
    if valor is None:
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    texto = unicodedata.normalize("NFKD", str(valor))
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", texto).strip().casefold()


def digitos_tarjeta(valor) -> str:
    if valor is None:
        return ""
    try:
        if pd.isna(valor):
            return ""
    except (TypeError, ValueError):
        pass
    return re.sub(r"\D", "", str(valor))


def extraer_ultimos_cinco(valor) -> str:
    digitos = digitos_tarjeta(valor)
    return digitos[-5:] if len(digitos) >= 5 else digitos


def parsear_fecha(valor):
    if valor is None:
        return None
    try:
        if pd.isna(valor):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(valor, pd.Timestamp):
        return valor.to_pydatetime()
    if isinstance(valor, datetime):
        return valor
    if isinstance(valor, date):
        return datetime.combine(valor, time.min)

    texto = str(valor).strip().replace("\u202f", " ").replace("\xa0", " ")
    texto = re.sub(r"\s+", " ", texto)
    try:
        return pd.to_datetime(texto, dayfirst=True, errors="raise").to_pydatetime()
    except Exception:
        return None


def parsear_hora(valor):
    if valor is None:
        return None
    if isinstance(valor, time):
        return valor.replace(second=0, microsecond=0)
    if isinstance(valor, datetime):
        return valor.time().replace(second=0, microsecond=0)
    if isinstance(valor, (int, float)) and not pd.isna(valor):
        segundos = int(round(float(valor) * 24 * 60 * 60)) % (24 * 60 * 60)
        return time(segundos // 3600, (segundos % 3600) // 60)
    if isinstance(valor, str):
        texto = valor.strip().upper().replace("\u202f", " ").replace("\xa0", " ")
        for fmt in ("%H:%M", "%H:%M:%S", "%I:%M %p", "%I:%M:%S %p", "%I:%M%p"):
            try:
                return datetime.strptime(texto, fmt).time().replace(second=0, microsecond=0)
            except ValueError:
                pass
    return None


def es_feriado_venezuela(fecha_valor: date, adicionales) -> bool:
    return (fecha_valor.month, fecha_valor.day) in FERIADOS_FIJOS or fecha_valor.isoformat() in adicionales


def clasificar_jornada(fecha_valor: date, adicionales) -> str:
    return "FER" if fecha_valor.weekday() >= 5 or es_feriado_venezuela(fecha_valor, adicionales) else "ORD"


def horas_entre(desde, hasta) -> float:
    desde_t = parsear_hora(desde)
    hasta_t = parsear_hora(hasta)
    if desde_t is None or hasta_t is None:
        return 0.0

    inicio = datetime.combine(date(2000, 1, 1), desde_t)
    fin = datetime.combine(date(2000, 1, 1), hasta_t)
    if fin < inicio:
        fin += timedelta(days=1)
    return max(0.0, (fin - inicio).total_seconds() / 3600)


def leer_bytes(archivo_obj) -> bytes:
    if hasattr(archivo_obj, "seek"):
        archivo_obj.seek(0)
        data = archivo_obj.read()
        archivo_obj.seek(0)
        return data
    elif isinstance(archivo_obj, (bytes, bytearray)):
        return bytes(archivo_obj)
    elif isinstance(archivo_obj, (str, Path)):
        return Path(archivo_obj).read_bytes()
    raise TypeError("Tipo de archivo no soportado para lectura en bytes.")


def cargar_dataframe_accesos(archivo) -> pd.DataFrame:
    contenido = leer_bytes(archivo)
    crudo = pd.read_excel(io.BytesIO(contenido), header=None, dtype=object)

    fila_encabezado = None
    for idx, fila in crudo.iterrows():
        valores = " ".join(normalizar_texto(v) for v in fila.tolist())
        if "fecha" in valores and ("personas" in valores or "trabajador" in valores or "nombre" in valores) and ("tarjeta" in valores or "carnet" in valores):
            fila_encabezado = idx
            break

    if fila_encabezado is None:
        for idx, fila in crudo.iterrows():
            valores = " ".join(normalizar_texto(v) for v in fila.tolist())
            if "fecha" in valores and ("tarjeta" in valores or "carnet" in valores):
                fila_encabezado = idx
                break

    if fila_encabezado is None:
        fila_encabezado = 0

    df = pd.read_excel(io.BytesIO(contenido), skiprows=fila_encabezado, dtype=object)
    df.columns = [str(c).strip().casefold() for c in df.columns]

    col_fecha = next((c for c in df.columns if "fecha" in c or "dia" in c), None)
    col_tarjeta = next((c for c in df.columns if "tarjeta" in c or "carnet" in c), None)
    col_persona = next((c for c in df.columns if "persona" in c or "trabajador" in c or "nombre" in c), None)

    if not col_fecha or not col_tarjeta or not col_persona:
        if len(df.columns) < 3:
            raise ValueError("No se reconocieron las columnas del Excel. Se requieren fecha, tarjeta/carnet y persona/trabajador.")
        df.columns = [f"col_{i}" for i in range(len(df.columns))]
        col_fecha, col_tarjeta, col_persona = df.columns[:3]

    df = df[df[col_fecha].notna()].copy()
    df["_fecha_dt"] = df[col_fecha].apply(parsear_fecha)
    df = df[df["_fecha_dt"].notna()].copy()
    df["_nombre_norm"] = df[col_persona].apply(normalizar_texto)
    df["_tarjeta_5"] = df[col_tarjeta].apply(extraer_ultimos_cinco)
    return df


def filtrar_accesos_persona(df, nombre_busqueda: str, carnet: str) -> pd.DataFrame:
    nombre_norm = normalizar_texto(nombre_busqueda)
    tokens = [t for t in nombre_norm.split() if t]
    carnet_norm = re.sub(r"\D", "", carnet)[-5:]

    coincidencia_nombre = df["_nombre_norm"].apply(lambda n: bool(tokens) and all(token in n.split() for token in tokens))
    coincidencia_carnet = df["_tarjeta_5"].eq(carnet_norm)

    filtrado = df[coincidencia_nombre & coincidencia_carnet].copy()
    if filtrado.empty:
        filtrado = df[coincidencia_carnet].copy()

    if filtrado.empty:
        raise ValueError("No se encontraron marcajes con los últimos cinco dígitos del carnet.")
    return filtrado


def leer_excel_accesos(archivo, nombre_busqueda: str, carnet: str, adicionales, solo_feriados: bool = True):
    df = cargar_dataframe_accesos(archivo)
    filtrado = filtrar_accesos_persona(df, nombre_busqueda, carnet)
    filtrado = filtrado.sort_values("_fecha_dt")

    registros = []
    for fecha_dia, grupo in filtrado.groupby(filtrado["_fecha_dt"].dt.date):
        jornada = clasificar_jornada(fecha_dia, adicionales)
        if solo_feriados and jornada != "FER":
            continue

        marcas = sorted(grupo["_fecha_dt"].tolist())
        if not marcas:
            continue

        entrada, salida = marcas[0], marcas[-1]
        registros.append({
            "id": uuid.uuid4().hex[:8],
            "fecha": fecha_dia,
            "jornada": jornada,
            "entrada": entrada.strftime("%H:%M"),
            "salida": salida.strftime("%H:%M"),
            "horas": round(horas_entre(entrada.time(), salida.time()), 2),
            "descripcion": "",
        })

    return registros, len(filtrado)


def buscar_horarios_en_excel(archivo, nombre_busqueda, carnet, fecha_buscada):
    try:
        df = cargar_dataframe_accesos(archivo)
        filtrado = filtrar_accesos_persona(df, nombre_busqueda, carnet)
        df_dia = filtrado[filtrado["_fecha_dt"].dt.date == fecha_buscada].sort_values("_fecha_dt")

        if not df_dia.empty:
            marcas = df_dia["_fecha_dt"].tolist()
            return (
                marcas[0].time().replace(second=0, microsecond=0),
                marcas[-1].time().replace(second=0, microsecond=0),
            )
    except Exception:
        pass
    return None, None


def validar_registros(df):
    errores = []
    vistos = set()
    for idx, r in df.iterrows():
        fila = idx + 1
        fecha = parsear_fecha(r.get("fecha"))
        entrada = parsear_hora(r.get("entrada"))
        salida = parsear_hora(r.get("salida"))
        jornada = str(r.get("jornada", "")).strip().upper()
        desc = str(r.get("descripcion", "") or "").strip()
        if fecha is None:
            errores.append(f"Fila {fila}: fecha inválida o vacía.")
        if entrada is None or salida is None:
            errores.append(f"Fila {fila}: hora de entrada o salida inválida.")
        elif entrada == salida:
            errores.append(f"Fila {fila}: entrada y salida son iguales.")
        if jornada not in {"ORD", "FER"}:
            errores.append(f"Fila {fila}: jornada debe ser ORD o FER.")
        if not desc:
            errores.append(f"Fila {fila}: falta la descripción.")
        if fecha and entrada and salida:
            clave = (fecha.date(), entrada, salida)
            if clave in vistos:
                errores.append(f"Fila {fila}: posible jornada duplicada.")
            vistos.add(clave)
    return errores


def encontrar_hoja_plantilla(wb):
    if NOMBRE_HOJA_PREFERIDA in wb.sheetnames:
        return wb[NOMBRE_HOJA_PREFERIDA]
    for ws in wb.worksheets:
        if ws["A5"].value and "AUTORIZACIÓN" in str(ws["A5"].value).upper():
            return ws
    return wb.active


def generar_excel_oficial(nombre: str, carnet: str, area: str, registros: list[dict], plantilla_path) -> bytes:
    if not os.path.exists(plantilla_path):
        raise ValueError(f"No se encontró el archivo de plantilla '{plantilla_path}' en la carpeta.")

    if len(registros) > MAX_JORNADAS_PLANTILLA:
        raise ValueError(f"La plantilla tiene espacio máximo para {MAX_JORNADAS_PLANTILLA} jornadas, pero seleccionaste {len(registros)}.")

    wb = openpyxl.load_workbook(plantilla_path)
    ws = encontrar_hoja_plantilla(wb)

    ws["E6"] = f" {nombre.strip().upper()}"
    ws["B8"] = carnet.strip()
    ws["D8"] = f"                       ÁREA/ DPTO                                                                      {area.strip().upper()}"

    for fila in range(FILA_INICIO_DATOS, FILA_FIN_DATOS + 1):
        for columna in range(1, 9):
            celda = ws.cell(fila, columna)
            if not isinstance(celda, MergedCell):
                celda.value = None

    registros_ordenados = sorted(registros, key=lambda r: (str(r.get("fecha", "")), str(r.get("entrada", ""))))
    red_font = Font(name="Calibri", size=11, color=COLOR_ROJO)

    for indice, reg in enumerate(registros_ordenados):
        fila = FILA_INICIO_DATOS + indice
        fecha_dt = parsear_fecha(reg.get("fecha"))
        desde = parsear_hora(reg.get("entrada"))
        hasta = parsear_hora(reg.get("salida"))
        jornada = str(reg.get("jornada", "FER")).upper()
        descripcion = str(reg.get("descripcion") or "").strip()

        c_fecha = ws.cell(fila, 1, value=fecha_dt)
        c_ord = ws.cell(fila, 2, value="X" if jornada == "ORD" else None)
        c_fer = ws.cell(fila, 3, value="X" if jornada == "FER" else None)
        c_desde = ws.cell(fila, 4, value=desde)
        c_hasta = ws.cell(fila, 5, value=hasta)
        c_desc = ws.cell(fila, 6, value=descripcion)

        c_ord.font = red_font
        c_fer.font = red_font
        c_desde.font = red_font
        c_hasta.font = red_font

        if jornada == "FER":
            c_tot = ws.cell(fila, 8, value=f"=+(E{fila}-D{fila})*24")
            c_tot.font = red_font
        else:
            c_tot = ws.cell(fila, 7, value=f"=+(E{fila}-D{fila})*24")
            c_tot.font = red_font

        c_fecha.number_format = "dd/mm/yyyy"
        c_desde.number_format = "hh:mm"
        c_hasta.number_format = "hh:mm"
        c_desc.alignment = Alignment(vertical="top", wrap_text=True)

    cell_total_footer = ws.cell(28, 8)
    if not cell_total_footer.value or "=" not in str(cell_total_footer.value):
        fila_fin = FILA_INICIO_DATOS + len(registros_ordenados) - 1
        cell_total_footer.value = f"=SUM(H13:H{fila_fin})"
    cell_total_footer.font = red_font

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# --- INTERFAZ STREAMLIT ---
st.title("Sistema gestión de horas extras")
st.caption("Carga registros de acceso, revisa jornadas, completa descripciones abajo y exporta la plantilla oficial Excel.")

with st.sidebar:
    st.header("1. Archivo de accesos")
    excel = st.file_uploader("Excel de registros", type=["xlsx", "xls"])

    st.header("2. Datos del funcionario")
    nombre = st.text_input("Nombre completo", placeholder="NOMBRE Y APELLIDO")
    carnet = st.text_input("Carnet / últimos dígitos", max_chars=12)
    area = st.text_input("Área / Departamento", value="DGIF")
    
    st.header("3. Feriados adicionales")
    feriados_txt = st.text_area("Una fecha por línea (AAAA-MM-DD)", placeholder="2026-10-12\n2026-12-08")
    adicionales = set()
    for line in feriados_txt.splitlines():
        try: adicionales.add(pd.to_datetime(line.strip()).date().isoformat())
        except Exception: pass

if "registros" not in st.session_state:
    st.session_state.registros = pd.DataFrame(columns=["id", "fecha", "jornada", "entrada", "salida", "horas", "descripcion"])

col1, col2 = st.columns([1, 1])
with col1:
    if st.button("Importar jornadas del Excel", type="primary", disabled=not (nombre.strip() and carnet.strip() and excel)):
        try:
            rows, _ = leer_excel_accesos(excel, nombre, carnet, adicionales, solo_feriados=True)
            st.session_state.registros = pd.DataFrame(rows, columns=["id", "fecha", "jornada", "entrada", "salida", "horas", "descripcion"])
            st.success(f"Se importaron {len(rows)} jornadas. Revisa y completa las descripciones abajo.")
        except Exception as e:
            st.error(str(e))
st.divider()
st.subheader("Listado Jornadas")
st.caption("Edita fechas, horas o tipo de jornada directamente en la tabla.")

if not st.session_state.registros.empty:
    edited = st.data_editor(
        st.session_state.registros, 
        key="editor_jornadas", 
        num_rows="dynamic", 
        use_container_width=True,
        column_config={
            "id": st.column_config.TextColumn("ID", disabled=True, width="small"),
            "fecha": st.column_config.DateColumn("Día", format="DD/MM/YYYY"),
            "jornada": st.column_config.SelectboxColumn("Jornada", options=["ORD", "FER"], required=True),
            "entrada": st.column_config.TextColumn("Desde", help="Ej. 08:00 AM"),
            "salida": st.column_config.TextColumn("Hasta", help="Ej. 01:00 PM"),
            "horas": st.column_config.NumberColumn("Horas", min_value=0.0, max_value=24.0, step=0.25, format="%.2f"),
            "descripcion": st.column_config.TextColumn("Descripción corta", width="medium"),
        }
    )
    edited = edited.copy()
    edited["horas"] = edited.apply(
        lambda r: round(horas_entre(parsear_hora(r.get("entrada")), parsear_hora(r.get("salida"))), 2)
        if parsear_hora(r.get("entrada")) and parsear_hora(r.get("salida")) else 0.0,
        axis=1,
    )
    st.session_state.registros = edited

        # --- SECCIÓN DE REDACCIÓN DE DESCRIPCIONES (ABAJO) ---
    st.divider()
    st.subheader("Redacta aqui las descripciones de cada jornada")
    st.caption("Selecciona un día en la lista desplegable y escribe su descripción en la parte inferior de forma fluida una por una.")

    opciones_jornadas = {
        f"{r['fecha']} ({r['jornada']} - {r['entrada']} a {r['salida']})": r["id"]
        for _, r in st.session_state.registros.iterrows()
    }

    seleccion = st.selectbox("Selecciona la fecha / jornada a documentar", list(opciones_jornadas.keys()), key="select_jornada_desc")
    id_sel = opciones_jornadas[seleccion]

    reg_actual = st.session_state.registros.loc[st.session_state.registros["id"] == id_sel].iloc[0]

    with st.form(f"form_desc_{id_sel}"):
        texto_ingresado = st.text_area(
            "Descripción del trabajo realizado para esta fecha",
            value=reg_actual["descripcion"],
            height=120,
            placeholder="Ej. Se brindó apoyo técnico en mantenimiento..."
        )
        aplicar_btn = st.form_submit_button("Guardar Descripción", type="primary")

        if aplicar_btn:
            st.session_state.registros.loc[st.session_state.registros["id"] == id_sel, "descripcion"] = texto_ingresado.strip()
            st.success("¡Descripción guardada exitosamente en el registro!")
            st.rerun()

    # --- CONSULTAR / AGREGAR DÍA ADICIONAL BUSCANDO EN EXCEL ---
    st.divider()
    st.subheader("Agregar día manualmente")
    with st.form("form_jornada_manual"):
        col_m1, col_m2 = st.columns([1.2, 1])
        fecha_manual = col_m1.date_input("Fecha a consultar/agregar", value=date.today(), format="DD/MM/YYYY")
        jornada_manual = col_m2.selectbox("Tipo de jornada", ["FER", "ORD"], index=0 if clasificar_jornada(date.today(), adicionales) == "FER" else 1)

        ent_auto, sal_auto = None, None
        if excel is not None and nombre.strip() and len(digitos_tarjeta(carnet)) >= 5:
            ent_auto, sal_auto = buscar_horarios_en_excel(excel, nombre, carnet, fecha_manual)

        col_m3, col_m4 = st.columns(2)
        desde_manual = col_m3.time_input("Hora desde (entrada)", value=ent_auto if ent_auto else time(8, 0), step=300)
        hasta_manual = col_m4.time_input("Hora hasta (salida)", value=sal_auto if sal_auto else time(14, 0), step=300)
        desc_manual = st.text_area("Descripción de la actividad para este día", placeholder="Describe las actividades realizadas...", height=90)

        if ent_auto and sal_auto:
            st.info(f"💡 Marcaje detectado en Excel para esta fecha: {ent_auto.strftime('%H:%M')} a {sal_auto.strftime('%H:%M')}")
        else:
            st.caption("No hay marcaje automático registrado en el Excel para esta fecha específica; puedes ajustarlo manualmente.")

        agregar_manual_btn = st.form_submit_button("Añadir / Actualizar Día en el Listado", type="primary")

        if agregar_manual_btn:
            if not desc_manual.strip():
                st.error("Debes ingresar una descripción.")
            else:
                fecha_str = fecha_manual
                # Reemplazar si ya existe la fecha
                st.session_state.registros = st.session_state.registros[st.session_state.registros["fecha"] != fecha_str]
                nuevo_reg = pd.DataFrame([{
                    "id": uuid.uuid4().hex[:8],
                    "fecha": fecha_str,
                    "jornada": jornada_manual,
                    "entrada": desde_manual.strftime("%H:%M"),
                    "salida": hasta_manual.strftime("%H:%M"),
                    "horas": round(horas_entre(desde_manual, hasta_manual), 2),
                    "descripcion": desc_manual.strip()
                }])
                st.session_state.registros = pd.concat([st.session_state.registros, nuevo_reg], ignore_index=True)
                st.success("¡Jornada agregada/actualizada con éxito!")
                st.rerun()

    # --- RESUMEN Y EXPORTACIÓN ---
    st.divider()
    st.subheader("Resumen")

    df = st.session_state.registros.copy()
    if "horas" in df:
        df["horas"] = pd.to_numeric(df["horas"], errors="coerce").fillna(0)
    total = float(df["horas"].sum()) if not df.empty else 0
    ord_total = float(df.loc[df["jornada"] == "ORD", "horas"].sum()) if not df.empty else 0
    fer_total = float(df.loc[df["jornada"] == "FER", "horas"].sum()) if not df.empty else 0

    m1, m2, m3 = st.columns(3)
    m1.metric("Total horas", f"{total:.2f}")
    m2.metric("Ordinarias (ORD)", f"{ord_total:.2f}")
    m3.metric("Feriados (FER)", f"{fer_total:.2f}")

    errores = validar_registros(df)
    if errores:
        with st.expander(f"Validaciones pendientes ({len(errores)})", expanded=True):
            for e in errores: st.warning(e)

    if st.button("Descargar Excel Rellenado", type="primary", disabled=bool(errores) or not nombre.strip() or not carnet.strip()):
        try:
            excel_bytes = generar_excel_oficial(
                nombre, 
                carnet, 
                area, 
                df.to_dict("records"), 
                PLANTILLA_DEFAULT_PATH
            )
            nombre_seguro = normalizar_texto(nombre).replace(" ", "_")
            st.download_button(
                label="Click para descargar",
                data=excel_bytes,
                file_name=f"Horas_Extras_{nombre_seguro}.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )
            st.success("¡Plantilla rellenada y lista para descargar!")
        except Exception as e:
            st.error(f"Error al generar la plantilla Excel: {e}")
else:
    st.info("Carga un archivo Excel de accesos o agrega una jornada manual para comenzar.")